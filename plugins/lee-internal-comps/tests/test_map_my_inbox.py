"""map-my-inbox helper (gi-plugins#176): the locked Slices shape, the voice lint,
the summary-only thread contract, and the stdlib xlsx writer.

Loaded by file path under a unique module name (conftest rule). openpyxl is a
TEST dependency only; the helper itself is stdlib.
"""

import importlib.util
import json
import os
import subprocess
import sys
import tempfile

import pytest
from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
SKILL_DIR = os.path.join(HERE, "..", "skills", "map-my-inbox")
MAP_PY = os.path.join(SKILL_DIR, "map.py")


def _load():
    name = "helpers_map_my_inbox_map"
    if name in sys.modules:
        return sys.modules[name]
    spec = importlib.util.spec_from_file_location(name, MAP_PY)
    mod = importlib.util.module_from_spec(spec)
    sys.modules[name] = mod
    spec.loader.exec_module(mod)
    return mod


def _row(**over):
    r = {
        "n": 1,
        "slice": "Prospect emails asking about space",
        "looks_like": "A new name asking about space in Wilmington. Short, one question.",
        "count": "6 (30 days)",
        "reply_or_feed": "Reply, they start it",
        "fits": "This is the front door of your deal flow: someone writes, you reply to set a call.",
        "needs": "What they are looking for, which is in the email. Whether you already talked, which is in your deal flow sheet.",
        "does": "New name: reply with two times for a call. Someone you know: reply and update their row.",
        "examples": [
            {"label": "Space in Wilmington?", "url": "https://outlook.office.com/mail/id/abc"},
            {"label": "Flex space question", "url": "https://outlook.office.com/mail/id/def"},
        ],
    }
    r.update(over)
    return r


THREAD = {"sender": "a@b.com", "subject": "Space?", "date": "2026-09-10", "count": 2, "line": "asks about flex space"}


def test_columns_are_the_locked_22():
    m = _load()
    assert len(m.SLICE_COLUMNS) == 22
    assert m.SLICE_COLUMNS[0] == "#"
    assert m.SLICE_COLUMNS[5] == "Where this fits in your work"
    assert m.SLICE_COLUMNS[6] == "What you'd need to know to answer it, and where that lives"
    assert m.SLICE_COLUMNS[7] == "What you do with it, and when"
    assert m.SLICE_COLUMNS[11] == "Real category? (yes / no / merge with #)"
    assert m.SLICE_COLUMNS[15] == "AI-ification stage"
    assert m.SLICE_COLUMNS[21] == "Progress notes"
    assert len(m.STAGES) == 9


def test_columns_match_the_worker_literal_when_the_lee_repo_is_present():
    """The Worker carries the same literal; the lee repo's check:parity pins it
    against gi-plugins main. When the sibling checkout is present locally, pin
    it here too so the drift is caught before a push."""
    worker = os.environ.get("GI_LEE_WORKER_INBOX_MAP")
    if not worker:
        # The GI parent repo root is the parent of the gi-plugins checkout, worktree or not.
        try:
            top = subprocess.run(["git", "rev-parse", "--show-toplevel"], cwd=HERE, capture_output=True, text=True).stdout.strip()
            gi_root = os.path.dirname(top) if os.path.basename(top) == "gi-plugins" else os.path.dirname(os.path.dirname(os.path.dirname(top)))
        except Exception:
            gi_root = ""
        worker = os.path.join(gi_root, "30_clients", "lee_and_associates", "sow_1_analyst_pilot", "mcp-server", "src", "tools", "leader_context", "inbox_map.ts")
    if not worker or not os.path.exists(worker):
        pytest.skip("lee repo not checked out beside gi-plugins (set GI_LEE_WORKER_INBOX_MAP)")
    src = open(worker, encoding="utf-8").read()
    start = src.index("export const SLICE_COLUMNS")
    end = src.index("];", start)
    theirs = [json.loads('"' + x + '"') for x in __import__("re").findall(r'"((?:[^"\\]|\\.)*)"', src[start:end])]
    assert theirs == _load().SLICE_COLUMNS


@pytest.mark.parametrize(
    "field,text,why",
    [
        ("fits", "Front door — of deal flow", "dash"),
        ("does", "call – or email", "dash"),
        ("needs", "Q2: the criteria", "Q1/Q2/Q3"),
        ("fits", "See card #482 for it", "card"),
        ("fits", "This is stage 3 of the funnel", "stage number"),
        ("looks_like", "See card #482", "card"),
        ("fits", "Tracked as lee#581", "repo card"),
        ("fits", "Part of L2: broker support", "chart process id"),
        ("fits", "Guess: the front door", "Guess"),
    ],
)
def test_voice_lint_refuses_internal_jargon(field, text, why):
    m = _load()
    probs = m.lint_slices([_row(**{field: text})])
    assert probs, f"expected a lint hit for {text!r}"
    assert why.lower() in probs[0].lower()


def test_voice_lint_passes_plain_words():
    m = _load()
    assert m.lint_slices([_row()]) == []


def test_validate_slices_rejects_unknown_key_and_missing_name_and_cap():
    m = _load()
    with pytest.raises(m.MapError):
        m.validate_slices([dict(_row(), stage="3")])
    with pytest.raises(m.MapError):
        m.validate_slices([_row(slice="")])
    with pytest.raises(m.MapError):
        m.validate_slices([_row(n=i) for i in range(41)])
    with pytest.raises(m.MapError):
        m.validate_slices([_row(examples=[{"label": "x", "url": "https://x"}] * 4)])


def test_validate_slices_mirrors_the_worker_caps_and_rejects_duplicate_names():
    m = _load()
    with pytest.raises(m.MapError):
        m.validate_slices([_row(count="")])
    with pytest.raises(m.MapError):
        m.validate_slices([_row(fits="x" * 2001)])
    with pytest.raises(m.MapError):
        m.validate_slices([_row(n=1, slice="Same"), _row(n=2, slice="same ")])
    # 40 rows is the cap, not over it
    assert len(m.validate_slices([_row(n=i, slice=f"Slice {i}") for i in range(40)])) == 40


def test_control_characters_and_markup_never_break_the_workbook(tmp_path):
    m = _load()
    rows = m.validate_slices([_row(
        slice="Tenants & owners <asking> \"now\"",
        examples=[{"label": "Re: rent roll\x0b & terms <Q3>", "url": "https://mail.example/?a=1&b=2"}],
    )])
    out = m.build_xlsx(rows, ["none"], str(tmp_path / "x.xlsx"))
    import zipfile as _z
    import xml.etree.ElementTree as ET
    with _z.ZipFile(out) as z:
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                ET.fromstring(z.read(name))
    wb = load_workbook(out)
    ws = wb["Slices"]
    assert ws["B2"].value == 'Tenants & owners <asking> "now"'
    assert ws["I2"].value == "Re: rent roll & terms <Q3>"
    assert ws["I2"].hyperlink.target == "https://mail.example/?a=1&b=2"


def test_missing_measurables_file_is_loud_unless_waived(tmp_path):
    m = _load()
    with pytest.raises(m.MapError):
        m.read_measurables(str(tmp_path / "nope.txt"))
    with pytest.raises(m.MapError):
        m.read_measurables(None)
    assert m.read_measurables(None, allow_missing=True) == ["none"]


def test_sent_rows_carry_the_recipient_in_to():
    m = _load()
    out = m.validate_threads([dict(THREAD, to="amy@york.example"), THREAD])
    assert out[0]["to"] == "amy@york.example"
    assert "to" not in out[1]


def test_voice_lint_tolerates_ordinary_cre_text():
    m = _load()
    assert m.lint_slices([_row(fits="Q3 rent roll requests for Suite 200 come in by email.")]) == []


def test_validate_threads_rejects_bodies_and_caps_the_line():
    m = _load()
    with pytest.raises(m.MapError) as e:
        m.validate_threads([dict(THREAD, body="Hi Sandy")])
    assert "never a body" in str(e.value)
    out = m.validate_threads([dict(THREAD, line="x" * 500)])
    assert len(out[0]["line"]) == m.MAX_LINE_CHARS


def test_select_threads_takes_up_to_ten_per_slice_and_caps_at_200():
    m = _load()
    slices = [_row(n=1, slice="A"), _row(n=2, slice="B")]
    threads = [dict(THREAD, slice="A") for _ in range(30)] + [dict(THREAD, slice="B") for _ in range(5)] + [dict(THREAD) for _ in range(400)]
    picked = m.select_threads(threads, slices)
    assert len(picked) == 200
    assert sum(1 for t in picked if t.get("slice") == "A") == 10
    assert sum(1 for t in picked if t.get("slice") == "B") == 5


def test_build_writes_a_one_tab_workbook_with_dropdowns_and_links(tmp_path):
    m = _load()
    slices = [m.validate_slices([_row(), _row(n=2, slice="Listing alerts", reply_or_feed="Feed, no reply", examples=[])])][0]
    out = m.build_xlsx(slices, ["Deals closed", "Live requirements", "none"], str(tmp_path / "inbox-map.xlsx"))
    wb = load_workbook(out)
    visible = [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]
    assert visible == ["Slices"]
    ws = wb["Slices"]
    headers = [ws.cell(row=1, column=c).value for c in range(1, 23)]
    assert headers == m.SLICE_COLUMNS
    assert ws["B2"].value == "Prospect emails asking about space"
    assert ws["I2"].value == "Space in Wilmington?"
    assert ws["I2"].hyperlink is not None and ws["I2"].hyperlink.target == "https://outlook.office.com/mail/id/abc"
    assert ws["L2"].value is None and ws["P2"].value is None
    sqrefs = [str(dv.sqref) for dv in ws.data_validations.dataValidation]
    assert any(s.startswith("L2") for s in sqrefs)
    assert any(s.startswith("M2") for s in sqrefs)
    assert any(s.startswith("N2") for s in sqrefs)
    assert any(s.startswith("P2") for s in sqrefs)
    lists = wb["Lists"]
    col_b = [lists.cell(row=r, column=2).value for r in range(1, 4)]
    assert col_b == ["Deals closed", "Live requirements", "none"]
    col_d = [lists.cell(row=r, column=4).value for r in range(1, 10)]
    assert col_d[5] == "Architecture spike"
    assert ws.freeze_panes == "C2"


def test_cli_build_prints_the_exact_submit_call(tmp_path):
    slices = tmp_path / "slices.json"
    threads = tmp_path / "threads.json"
    meas = tmp_path / "measurables.txt"
    slices.write_text(json.dumps([_row()]), encoding="utf-8")
    threads.write_text(json.dumps([dict(THREAD, slice="Prospect emails asking about space")]), encoding="utf-8")
    meas.write_text("Deals closed\n", encoding="utf-8")
    p = subprocess.run(
        [sys.executable, MAP_PY, "build", str(slices), str(threads), "--out", str(tmp_path), "--measurables", str(meas), "--date", "2026-09-12"],
        capture_output=True, text=True,
    )
    assert p.returncode == 0, p.stdout + p.stderr
    assert "WROTE" in p.stdout and "inbox-map-2026-09-12.xlsx" in p.stdout
    assert (tmp_path / "inbox-map-2026-09-12.xlsx").exists()
    call = json.loads(p.stdout[p.stdout.index("{"):])
    assert call["tool"] == "submit_inbox_map"
    assert call["args"]["slices"][0]["slice"] == "Prospect emails asking about space"
    assert call["args"]["threads"][0]["line"] == "asks about flex space"
    assert "body" not in call["args"]["threads"][0]


def test_cli_lint_fails_loudly_on_a_dash(tmp_path):
    slices = tmp_path / "slices.json"
    slices.write_text(json.dumps([_row(fits="the door — open")]), encoding="utf-8")
    p = subprocess.run([sys.executable, MAP_PY, "lint", str(slices)], capture_output=True, text=True)
    assert p.returncode == 1
    assert "VOICE LINT FAILED" in p.stdout


def test_skill_md_never_reuses_a_prior_map_and_never_builds():
    """Run 4 (2026-09-13) reused a six-hour-old map and drifted into slice building. The
    SKILL.md must carry the fresh-map rule and must not instruct a read of the inbox_map kind."""
    text = open(os.path.join(SKILL_DIR, "SKILL.md"), encoding="utf-8").read()
    assert "Every run builds a fresh map" in text
    assert "never ask\nwhether to reuse it" in text or "never ask whether to reuse it" in text.replace("\n", " ")
    assert "This run maps; it never builds" in text
    assert 'kind: "inbox_map"' not in text.replace("Do not call `get_my_context` with `kind: \"inbox_map\"`", "")
